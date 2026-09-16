"""Тесты сборки .docx из Protocol."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from app.docx.builder import build_docx, save_docx
from app.protocol.constants import PLACEHOLDER
from app.protocol.models import Protocol, Task


# ------------------------------------------------------------------ fixtures


def _make_full_protocol() -> Protocol:
    return Protocol(
        title="Проект Альфа",
        date="2024-09-12",
        participants=["Петя", "Маша"],
        discussion=["Обсудили релиз.", "Обсудили бюджет."],
        decisions=["Перенести релиз."],
        tasks=[
            Task(text="Написать тесты", assignee="Маша", due_date="2024-09-20"),
            Task(text="Подготовить доки", assignee="Петя", due_date=PLACEHOLDER),
        ],
    )


def _make_empty_protocol() -> Protocol:
    return Protocol(
        title=PLACEHOLDER,
        date=PLACEHOLDER,
        participants=[],
        discussion=[],
        decisions=[],
        tasks=[],
    )


# ------------------------------------------------------------------ structure


def test_build_docx_returns_document() -> None:
    doc = build_docx(_make_full_protocol())
    assert doc is not None
    assert len(doc.paragraphs) > 0


def test_h1_contains_title() -> None:
    doc = build_docx(_make_full_protocol())
    h1 = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert len(h1) == 1
    assert "Проект Альфа" in h1[0].text


def test_h2_sections_present() -> None:
    doc = build_docx(_make_full_protocol())
    h2_texts = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert h2_texts == ["Обсуждение", "Решения", "Задачи"]


def test_meta_lines_present() -> None:
    doc = build_docx(_make_full_protocol())
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Дата:" in all_text
    assert "2024-09-12" in all_text
    assert "Участники:" in all_text
    assert "Петя, Маша" in all_text


def test_discussion_bullets_present() -> None:
    doc = build_docx(_make_full_protocol())
    bullet_texts = [
        p.text for p in doc.paragraphs
        if p.style.name == "List Bullet"
    ]
    assert "Обсудили релиз." in bullet_texts
    assert "Перенести релиз." in bullet_texts


# ------------------------------------------------------------------ tasks table


def test_tasks_table_has_correct_shape() -> None:
    """Таблица: 1 шапка + 2 задачи = 3 строки, 3 колонки."""
    doc = build_docx(_make_full_protocol())
    assert len(doc.tables) == 1

    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 3


def test_tasks_table_header() -> None:
    doc = build_docx(_make_full_protocol())
    header = doc.tables[0].rows[0].cells
    assert header[0].text == "Задача"
    assert header[1].text == "Ответственный"
    assert header[2].text == "Срок"


def test_tasks_content() -> None:
    doc = build_docx(_make_full_protocol())
    rows = doc.tables[0].rows
    assert rows[1].cells[0].text == "Написать тесты"
    assert rows[1].cells[1].text == "Маша"
    assert rows[1].cells[2].text == "2024-09-20"
    assert rows[2].cells[2].text == PLACEHOLDER


# ------------------------------------------------------------------ placeholders


def test_placeholder_run_is_italic_and_gray() -> None:
    """Ячейка с плейсхолдером должна иметь курсивный серый run."""
    doc = build_docx(_make_full_protocol())
    row = doc.tables[0].rows[2]  # задача с PLACEHOLDER в due_date
    cell = row.cells[2]

    placeholder_runs = [
        r for p in cell.paragraphs
        for r in p.runs
        if r.text == PLACEHOLDER
    ]
    assert placeholder_runs, "нет run-а с плейсхолдером"
    run = placeholder_runs[0]
    assert run.italic is True
    assert run.font.color.rgb is not None


def test_empty_protocol_renders_with_placeholders() -> None:
    """Пустой Protocol не падает, всё залито плейсхолдерами."""
    doc = build_docx(_make_empty_protocol())

    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert PLACEHOLDER in all_text  # в H1, дате, участниках

    # Таблица всё равно есть, одна строка данных
    table = doc.tables[0]
    assert len(table.rows) == 2  # шапка + 1 плейсхолдер
    data_row = table.rows[1]
    for cell in data_row.cells:
        assert cell.text == PLACEHOLDER


# ------------------------------------------------------------------ save


def test_save_docx_creates_file(tmp_path: Path) -> None:
    out = tmp_path / "nested" / "protocol.docx"
    save_docx(_make_full_protocol(), out)
    assert out.exists()
    assert out.stat().st_size > 0


def test_saved_file_is_readable_by_python_docx(tmp_path: Path) -> None:
    """Круг: save → load → проверили содержимое."""
    out = tmp_path / "protocol.docx"
    save_docx(_make_full_protocol(), out)

    doc = Document(str(out))
    h1 = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert h1 and "Проект Альфа" in h1[0].text